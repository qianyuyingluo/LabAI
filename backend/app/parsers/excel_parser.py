from pathlib import Path

from app.parsers.text_parser import format_cell


def parse_xlsx(path: Path) -> str:
    try:
        from openpyxl import load_workbook
    except ImportError:
        return "Install openpyxl to extract Excel workbook content."

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        lines: list[str] = [f"Workbook sheets: {', '.join(workbook.sheetnames)}"]
        for sheet in workbook.worksheets:
            lines.append(f"\nSheet: {sheet.title} ({sheet.max_row} rows x {sheet.max_column} columns)")
            for row in sheet.iter_rows(values_only=True):
                values = [format_cell(value) for value in row]
                if any(values):
                    lines.append(" | ".join(values))
        return "\n".join(lines)
    finally:
        workbook.close()
